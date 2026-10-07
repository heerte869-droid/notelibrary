from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pdfencrypt import StandardEncryption
from PIL import Image,ImageDraw,ImageFont
import zipfile,io
out=Path(__file__).resolve().parent/'Fixtures';out.mkdir(exist_ok=True)
(out/'lesson.md').write_text('# 水循环\n\n蒸发 → 凝结 → 降水。\n\n|阶段|含义|\n|---|---|\n|蒸发|液体变成气体|\n\n```python\nprint("保留代码")\n```\n',encoding='utf-8')
(out/'table.csv').write_bytes('科目,知识点\n物理,动量守恒\n'.encode('utf-16'))
(out/'gb18030.txt').write_bytes('中文旧编码：细胞与光合作用。'.encode('gb18030'))
(out/'page.html').write_text('<html><head><script>DO_NOT_EXECUTE</script></head><body><h1>太阳能</h1><p>光合作用 &amp; 叶绿体</p><img src="https://invalid.example/tracker" /></body></html>')
doc=Document();doc.add_heading('水循环学习资料',0);doc.add_paragraph('第一部分：蒸发、凝结与降水。');t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='阶段';t.cell(0,1).text='解释';t.cell(1,0).text='凝结';t.cell(1,1).text='水蒸气形成水滴';doc.save(out/'lesson.docx')
p=Presentation()
for title,body,note in [('第一部分','Water evaporates from the surface.','NOTE_ALPHA'),('第二部分','Vapor condenses into clouds.','NOTE_BETA'),('第三部分','Rain returns water to the ground.','NOTE_GAMMA')]:
 s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text=title;s.placeholders[1].text=body;s.notes_slide.notes_text_frame.text=note
# Intentional order differs from slide part filenames.
ids=p.slides._sldIdLst;last=ids[-1];ids.remove(last);ids.insert(0,last);p.save(out/'lesson.pptx')
blank=Presentation();blank.slides.add_slide(blank.slide_layouts[6]);blank.save(out/'blank.pptx')
w=Workbook();s=w.active;s.title='物理';s['A1']='数量';s['C1']='公式';s['A2']=7;s['C2']='=A2*2';s['D3']=True;w.create_sheet('化学')['B4']='原子';w.save(out/'lesson.xlsx')
c=Canvas(str(out/'lesson.pdf'))
for i,txt in enumerate(['Evaporation: liquid water becomes vapor.','Condensation: vapor forms droplets.','Precipitation: water falls as rain.']):
 c.setFont('Helvetica',18);c.drawString(50,760,f'Water Cycle - Page {i+1}');c.setFont('Helvetica',14);c.drawString(50,700,txt);c.showPage()
c.save()
c=Canvas(str(out/'locked.pdf'),encrypt=StandardEncryption('test-only-password',canPrint=0));c.drawString(50,700,'Locked sample');c.save()
image=Image.new('RGB',(1500,900),'white');d=ImageDraw.Draw(image);font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',65);d.text((80,150),'WATER CYCLE',fill='black',font=font);d.text((80,310),'Evaporation and rain',fill='black',font=font);image.save(out/'scanned.pdf','PDF',resolution=150)
# Real ZIP fixtures for rejected malformed paths and entities.
with zipfile.ZipFile(out/'unsafe.docx','w',zipfile.ZIP_DEFLATED) as z:z.writestr('../escape.txt','not extracted');z.writestr('word/document.xml','<document/>')
with zipfile.ZipFile(out/'entity.docx','w',zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml','<!DOCTYPE document [<!ENTITY x "expanded">]><document><p><t>&x;</t></p></document>')
print('Generated document fixtures:',len(list(out.iterdir())))
